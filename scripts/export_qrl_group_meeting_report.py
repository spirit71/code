#!/usr/bin/env python3
"""导出 QRL-BoQ 组会汇报的 CSV、XLSX 和 PPTX。"""

from __future__ import annotations

import csv
import json
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "2026-07-31_QRL-BoQ组会汇报"
DEPS = Path("/root/qrl_report_deps")
sys.path.insert(0, str(DEPS))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import xlsxwriter


E_MATRIX = [
    ["Pitts30k-test", 92.958, 93.339, 93.163, 93.090, 92.796],
    ["Nordland", 87.852, 87.141, 86.782, 86.862, 88.102],
    ["SPED", 91.104, 91.433, 91.104, 91.598, 91.433],
    ["AmsterTime", 63.444, 64.338, 63.363, 63.119, 64.338],
    ["Tokyo247", 96.508, 96.825, 96.825, 96.825, 96.825],
    ["SVOX-all", 98.567, 98.615, 98.610, 98.594, 98.621],
    ["Macro", 88.405, 88.615, 88.308, 88.348, 88.686],
]


def jload(path):
    with open(path, encoding="utf-8") as handle:
        return json.load(handle)


def write_csv(path, rows, headers):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def experiment_rows():
    return [
        {
            "阶段": "原始方法", "实验": "QRL可学习性", "名称": "Cross-view Query Reliability Learning",
            "研究动机": "验证单图O2内容和attention统计能否预测跨视角repeatability",
            "具体实现": "O2方向512+norm+entropy+max-attention=515维；MLP 515-128-1；rank伪标签",
            "核心结果": "Spearman 0.0363→0.4415；reliability std 0.00008→0.12525；gate 0.8161~1.1993",
            "判断": "可学习但不等于检索有效", "证据路径": "src/query_reliability.py; src/model.py; 2026-07-30_QRL-BoQ_方法代码实验结果与下一步完整报告.md",
        },
        {
            "阶段": "因果矩阵", "实验": "E0", "名称": "Original BoQ baseline",
            "研究动机": "建立不含QRL监督和gate的严格基线",
            "具体实现": "原始BoQ；O1/O2→FC→8192维descriptor",
            "核心结果": "六集Macro R@1=88.4054", "判断": "后续所有方法的主对照",
            "证据路径": "/root/qrl_e_matrix/analysis/qrl_e_matrix_results.csv",
        },
        {
            "阶段": "因果矩阵", "实验": "E1", "名称": "Auxiliary reliability supervision, gate off",
            "研究动机": "隔离QRL辅助监督及训练轨迹是否有价值",
            "具体实现": "learned QRL+rank loss；reliability_gate_mode=none",
            "核心结果": "Macro R@1=88.6154；相对E0 +0.2100pp",
            "判断": "弱Top-1信号，高K/逐域不一致", "证据路径": "/root/qrl_e_matrix/analysis/qrl_e_matrix_results.csv",
        },
        {
            "阶段": "因果矩阵", "实验": "E2", "名称": "Full learned QRL + centered-residual gate",
            "研究动机": "验证预测可靠性软加权O2能否提升地点检索",
            "具体实现": "w=clip(1+0.5(r-mean(r)),0.5,1.5)；只加权O2",
            "核心结果": "Macro R@1=88.3080；E2-E0=-0.0974pp；R@10=-0.2169pp；R@20=-0.2586pp",
            "判断": "完整QRL没有稳定提升", "证据路径": "/root/qrl_e_matrix/analysis/qrl_e_matrix_results.csv",
        },
        {
            "阶段": "因果矩阵", "实验": "E2-gateoff", "名称": "Same-weight gate-off control",
            "研究动机": "同checkpoint隔离gate的直接推理效应",
            "具体实现": "严格加载E2同一SHA；推理gate设为none",
            "核心结果": "Macro R@1=88.3481；E2-E2-gateoff=-0.0401pp",
            "判断": "gate无稳定直接收益", "证据路径": "/root/qrl_e_matrix/analysis/qrl_e_matrix_results.csv",
        },
        {
            "阶段": "因果矩阵", "实验": "E3", "名称": "Entropy reliability heuristic",
            "研究动机": "检查无需学习的attention熵是否足够",
            "具体实现": "r=1-normalized entropy",
            "核心结果": "Macro R@1=88.6859；E3-E0=+0.2805pp",
            "判断": "弱启发式Top-1信号，非全面Recall提升", "证据路径": "/root/qrl_e_matrix/analysis/qrl_e_matrix_results.csv",
        },
        {
            "阶段": "贡献诊断", "实验": "LOO-A/B", "名称": "O2 leave-one-query-out oracle",
            "研究动机": "验证64个O2 query是否具有不同的实例级检索贡献",
            "具体实现": "只mask query侧单个O2；reference gallery不变；计算正负margin变化",
            "核心结果": "Amster E0: baseline63.44/oracle-low50 84.57/random50 57.69；SPED E0:91.10/97.36/90.63",
            "判断": "贡献异质性真实存在，但oracle不可部署", "证据路径": "/root/qrl_retrieval_contribution_v2; 2026-07-30_QRL-BoQ_第二轮实验实施前审计与执行方案.md",
        },
        {
            "阶段": "机制诊断", "实验": "H1-H4", "名称": "Contribution mechanism analysis",
            "研究动机": "判断有害query来自冗余、图像难度、固定slot还是困难负样本混淆",
            "具体实现": "冗余Spearman、margin-有害率、正确/错误组、有益有害sign-switch分解",
            "核心结果": "冗余相关-0.0662~0.0003；margin相关-0.7971~-0.8740；sign-switch 89.06%~100%；负样本主导87.66%~95.96%",
            "判断": "贡献是instance-specific且主要由困难负样本混淆决定",
            "证据路径": "/root/qrl_utility_mechanisms/analysis_v2; 2026-07-30_QRL-BoQ_第三轮机制实验实施前审计与执行方案.md",
        },
        {
            "阶段": "标签升级", "实验": "Hard-label", "名称": "Global hard-negative utility labels",
            "研究动机": "把target从同地点稳定性推进到正样本-全图库困难负样本margin贡献",
            "具体实现": "GSV global hard negatives训练；MSLS-val全gallery验证；LOO contribution",
            "核心结果": "MSLS E0 R@1=92.5676%；negative contribution=28.18%；projection cosine≥0.9999994",
            "判断": "标签更接近检索，但单query输入可能信息不足", "证据路径": "/root/qrl_hard_negative_labels/summary.json",
        },
        {
            "阶段": "效用预测", "实验": "U0", "名称": "Query-only utility predictor",
            "研究动机": "测试单个O2向量能否预测其检索贡献",
            "具体实现": "冻结BoQ；共享MLP 512-128-1；3 seeds；pairwise ranking loss",
            "核心结果": "Spearman 0.0620 vs static 0.1433；pair 0.5322 vs 0.5744；per-index 0.0256",
            "判断": "单query内容不足，低于static", "证据路径": "/root/qrl_u0_global_hard_msls/summary.json",
        },
        {
            "阶段": "效用预测", "实验": "U1", "名称": "Image-context utility predictor",
            "研究动机": "检验同图64-query上下文能否补足U0信息",
            "具体实现": "输入[q,mean(q),q*mean,q-mean]，2048维→256→1",
            "核心结果": "Spearman 0.0651；pair 0.5341；per-index -0.0520；仍低于static",
            "判断": "简单均值上下文无效", "证据路径": "/root/qrl_u1_global_hard_msls/summary.json",
        },
        {
            "阶段": "候选感知", "实验": "C0", "名称": "Candidate-aware utility predictor",
            "研究动机": "贡献是关系量，显式加入Top-20 candidate信息",
            "具体实现": "O2 512+六个candidate统计=518维→128→1",
            "核心结果": "Spearman 0.0791；pair 0.5406；per-index 0.1963",
            "判断": "实例差异可预测性明显提高，但仍低于static总体排序", "证据路径": "/root/qrl_c0_candidate_aware_msls/summary.json",
        },
        {
            "阶段": "候选感知", "实验": "C1", "名称": "Static prior + candidate residual",
            "研究动机": "保留稳定slot先验，同时学习图像/候选特定残差",
            "具体实现": "score=γ·static_prior+MLP([O2,candidate_stats])；residual零初始化",
            "核心结果": "Spearman 0.1495>0.1433；pair 0.5775>0.5744；per-index 0.2186",
            "判断": "唯一稳定超过static的预测机制，但增幅小", "证据路径": "/root/qrl_c1_static_candidate_residual_msls/summary.json",
        },
        {
            "阶段": "检索闭环", "实验": "C2", "名称": "Centered bounded soft gate",
            "研究动机": "把C1分数转换为温和O2权重",
            "具体实现": "zscore→tanh→中心化；w=1+β·residual；β网格",
            "核心结果": "MSLS最佳β=0；685/712/717/719",
            "判断": "关闭gate最佳", "证据路径": "/root/qrl_c2_soft_gate_msls/summary.json",
        },
        {
            "阶段": "检索闭环", "实验": "C3", "名称": "Hard removal",
            "研究动机": "直接删除预测最低贡献query，检验排序因果性",
            "具体实现": "query侧删除最低1/2/4/8；reference不变；重新FC投影",
            "核心结果": "MSLS remove2=684/714/715/720；六测试R@1 0胜1平5负，macro -0.1721pp",
            "判断": "低预测分数不等于可安全删除", "证据路径": "/root/qrl_c3_hard_removal_msls/summary.json; /root/qrl_c0_c10_test_diagnostic/all_results.csv",
        },
        {
            "阶段": "检索闭环", "实验": "C4", "名称": "Cross-slot max Top-20 reranking",
            "研究动机": "绕过FC混合，直接用O2集合关系重排候选",
            "具体实现": "query slot对candidate 64 slots取max；C1 rank weighting；global_z+α local_z",
            "核心结果": "MSLS α=.2:685/713/715/719；六测试2胜1平3负，macro -0.0269pp",
            "判断": "SPED正向但跨域不稳定", "证据路径": "/root/qrl_c4_top20_reranking_msls/summary.json; /root/qrl_c0_c10_test_diagnostic/all_results.csv",
        },
        {
            "阶段": "检索闭环", "实验": "C5", "名称": "Same-slot reranking",
            "研究动机": "避免cross-slot max过激，测试固定query语义对齐",
            "具体实现": "只比较同编号slot cosine；Top-20融合",
            "核心结果": "MSLS最佳α=0，等于E0",
            "判断": "固定slot对应不足", "证据路径": "/root/qrl_c5_same_slot_reranking_msls/summary.json",
        },
        {
            "阶段": "检索闭环", "实验": "C6", "名称": "Chamfer/MNN bidirectional reranking",
            "研究动机": "用双向集合匹配降低单向max误配",
            "具体实现": "symmetric Chamfer与mutual nearest neighbor；C1/static/uniform",
            "核心结果": "平均mutual matches=15.36；最佳α=0",
            "判断": "更复杂匹配仍无Recall收益", "证据路径": "/root/qrl_c6_bidirectional_reranking_msls/summary.json",
        },
        {
            "阶段": "学习重排", "实验": "C7", "名称": "GSV supervised Candidate RankNet",
            "研究动机": "让模型学习global/local特征的融合而非手工alpha",
            "具体实现": "7特征；score=global+MLP(7→16→1)；GSV正负pairwise loss；3 seeds",
            "核心结果": "GSV Top20正样本覆盖99.8806%；MSLS full7=683/713/715/719；六测试macro -0.0631pp",
            "判断": "GSV候选过易，出现负迁移", "证据路径": "/root/qrl_c7_candidate_ranknet_msls/summary.json",
        },
        {
            "阶段": "学习重排", "实验": "C8", "名称": "MSLS cross-city RankNet, 1 epoch",
            "研究动机": "减少GSV→MSLS域差异并保持城市隔离",
            "具体实现": "CPH训练→SF验证、SF训练→CPH验证；1 epoch；3 seeds",
            "核心结果": "MSLS=685/712/717/719；六测试macro -0.0024pp，1胜4平1负",
            "判断": "近恒等映射，无明确增益", "证据路径": "/root/qrl_c8_msls_cross_city_ranknet/summary.json",
        },
        {
            "阶段": "学习重排", "实验": "C9", "名称": "MSLS cross-city RankNet, 20 epochs",
            "研究动机": "排除C8优化步数不足",
            "具体实现": "相同cross-city协议，训练20 epochs",
            "核心结果": "MSLS=683/713/715/719；六测试macro -0.0511pp，1胜1平4负",
            "判断": "充分优化后负迁移更明显", "证据路径": "/root/qrl_c9_msls_cross_city_ranknet_e20/summary.json",
        },
        {
            "阶段": "学习重排", "实验": "C10", "名称": "Global ambiguity routing",
            "研究动机": "只对Top1-Top2 gap小的困难query启用重排",
            "具体实现": "coverage 0/5/10/20/30/40/50%；gap排序路由",
            "核心结果": "MSLS最佳coverage=0；无query被路由",
            "判断": "当前RankNet不值得启用", "证据路径": "/root/qrl_c10_ambiguity_routing_msls/summary.json",
        },
        {
            "阶段": "跨域诊断", "实验": "Six-test", "名称": "Frozen six-benchmark diagnostic",
            "研究动机": "检验MSLS-val结论是否因域差异失真",
            "具体实现": "固定MSLS选择；六集322输入；L2 Top-20；不使用test GT调参",
            "核心结果": "C3/C4/C7/C8/C9 macro ΔR@1=-0.1721/-0.0269/-0.0631/-0.0024/-0.0511pp",
            "判断": "没有稳定跨域提升", "证据路径": "/root/qrl_c0_c10_test_diagnostic/all_results.csv",
        },
    ]


def test_rows():
    raw = list(csv.DictReader(open(
        "/root/qrl_c0_c10_test_diagnostic/all_results.csv", encoding="utf-8"
    )))
    by_dataset = {}
    for row in raw:
        by_dataset.setdefault(row["dataset"], {})[row["method"]] = row
    output = []
    for dataset, methods in by_dataset.items():
        base = methods["E0"]
        for method, row in methods.items():
            item = dict(row)
            for k in (1, 5, 10, 20):
                item[f"delta_r@{k}_pp"] = (
                    float(row[f"r@{k}"]) - float(base[f"r@{k}"])
                ) * 100
                item[f"r@{k}_pct"] = float(row[f"r@{k}"]) * 100
            output.append(item)
    return output


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(.55), Inches(.28), Inches(12.2), Inches(.65))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(25)
    p.font.bold = True
    p.font.color.rgb = RGBColor(16, 45, 74)
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(.6), Inches(.9), Inches(12), Inches(.35))
        p = sub.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(11)
        p.font.color.rgb = RGBColor(90, 100, 110)


def add_bullets(slide, items, x=.7, y=1.3, w=12, h=5.5, font=18):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p.text = text
        p.level = level
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(font - level * 2)
        p.space_after = Pt(9)
        p.font.color.rgb = RGBColor(32, 42, 52)


def add_table(slide, headers, rows, x=.4, y=1.25, w=12.5, h=5.5, font=10):
    table = slide.shapes.add_table(
        len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h)
    ).table
    for c, header in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = str(header)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(24, 78, 119)
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(value)
            if r % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(235, 243, 248)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.text_frame.paragraphs:
                p.font.name = "Microsoft YaHei"
                p.font.size = Pt(font)
                if row == table.rows[0]:
                    p.font.bold = True
                    p.font.color.rgb = RGBColor(255, 255, 255)
                p.alignment = PP_ALIGN.CENTER
    return table


def build_ppt(rows, tests):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    add_title(slide, "QRL-BoQ：从跨视角可靠性到检索贡献", "组会汇报｜2026-07-31｜所有结论均回溯当前代码、summary.json 与原始 CSV")
    add_bullets(slide, [
        "核心问题：64个query贡献是否不同？能否从单图预测？能否转化为Recall提升？",
        "研究路径：QRL可学习性 → 因果矩阵 → LOO贡献诊断 → U0/U1/C0/C1 → C2-C10检索闭环 → 六测试集",
        "最终结论：贡献异质性真实存在，但当前预测分数到检索收益的映射不稳定。",
    ], y=1.65, font=22)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "1. 总体研究逻辑")
    add_table(slide, ["层次", "问题", "代表实验", "结论"], [
        ["现象", "query贡献是否相同", "LOO-A/B", "显著异质，oracle空间大"],
        ["可预测性", "单图能否预测贡献", "U0/U1/C0/C1", "候选关系有帮助，C1略超static"],
        ["因果闭环", "预测分数能否提升检索", "C2-C10", "MSLS均未过门槛"],
        ["泛化", "是否只是MSLS域问题", "六测试集", "正负混合，无稳定增益"],
    ], font=15)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "2. 原始BoQ与QRL模块", "代码：src/boq.py:94-150；src/query_reliability.py:47-180")
    add_bullets(slide, [
        "BoQ：O1/O2均为[B,64,512]；拼接为[B,128,512]；FC 128→16；输出8192维L2归一化descriptor。",
        "QRL输入：normalize(O2) 512维 + query norm + attention entropy + max attention = [B,64,515]。",
        "QRL Head：共享MLP 515→128→1，sigmoid输出reliability [B,64]；它是排序分数，不是校准概率。",
        "centered-residual gate：w=clip(1+0.5(r−mean(r)),0.5,1.5)，只广播乘到O2。",
        "训练：L=L_MS+λ_rel L_rel+均值/方差正则；跨视角target stop-gradient。",
    ], font=18)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "3. QRL首先验证了什么？")
    add_table(slide, ["指标", "初期", "训练后", "含义"], [
        ["reliability-target Spearman", "0.0363", "0.4415", "跨视角排序可学习"],
        ["reliability std", "0.00008", "0.12525", "输出未塌缩"],
        ["gate mean", "1.0000", "1.0000", "中心化保持整体尺度"],
        ["gate min/max", "~1/~1", "0.8161/1.1993", "产生温和动态加权"],
    ], font=14)
    add_bullets(slide, ["边界：学会repeatability排序 ≠ 学会retrieval contribution ≠ Recall提升。"], y=6.25, h=.5, font=17)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "4. E0/E1/E2/E2-gateoff/E3 因果矩阵")
    add_table(slide, ["数据集", "E0", "E1", "E2", "E2-off", "E3"], E_MATRIX[:-1], font=11)
    add_bullets(slide, [
        "Macro R@1：E0 88.405；E1 88.615；E2 88.308；E2-off 88.348；E3 88.686。",
        "最干净gate效应：E2−E2-off = −0.0401 pp；完整E2−E0 = −0.0974 pp。",
    ], y=6.15, h=.8, font=15)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "5. 为什么转向LOO检索贡献？")
    add_bullets(slide, [
        "原target回答“同地点其他视角能否重现”，但检索还要求“不要匹配困难负地点”。",
        "LOO contribution：依次移除每个O2 query，观察正样本与最难负样本margin变化。",
        "reference gallery保持原始descriptor，只改变query侧，避免不公平地同时修改图库。",
        "oracle使用测试GT，仅用于回答‘贡献差异是否存在’，不可作为部署性能。",
    ], font=19)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "6. LOO贡献诊断：异质性确实存在")
    add_table(slide, ["数据集/模型", "Baseline R@1", "Random50", "Oracle-low50", "Oracle-high50", "负贡献比例"], [
        ["Amster E0", "63.44", "57.69", "84.57", "16.82", "39.51%"],
        ["Amster E2-off", "63.12", "58.36", "83.75", "16.90", "40.16%"],
        ["Amster E2-on", "63.36", "57.99", "84.24", "14.38", "40.42%"],
        ["SPED E0", "91.10", "90.63", "97.36", "67.38", "18.87%"],
        ["SPED E2-off", "91.60", "91.04", "96.21", "66.56", "19.01%"],
        ["SPED E2-on", "91.10", "90.94", "96.71", "63.26", "20.47%"],
    ], font=11)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "7. 机制诊断：不是简单的冗余问题")
    add_table(slide, ["证据", "数值", "解释"], [
        ["冗余度-贡献图内Spearman", "-0.0662~0.0003", "相似query不必然无用"],
        ["margin-有害query比例", "-0.7971~-0.8740", "困难图像积累更多有害query"],
        ["固定slot sign-switch", "89.06%~100%", "贡献随图像变化，非固定slot属性"],
        ["困难负样本项主导", "87.66%~95.96%", "主要问题是跨地点混淆"],
        ["错误/正确图有害率", "44~49% vs 8~22%", "错误图内部确有更多干扰query"],
    ], font=14)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "8. U0/U1：单图内容和简单上下文不够")
    add_table(slide, ["实验", "输入", "图内Spearman", "Pair Acc.", "Per-index", "Static"], [
        ["U0", "O2:512", "0.0620", "0.5322", "0.0256", "0.1433/0.5744"],
        ["U1", "q,mean,q×mean,q−mean", "0.0651", "0.5341", "-0.0520", "0.1433/0.5744"],
    ], font=14)
    add_bullets(slide, ["结论：utility是query、正样本和竞争负样本共同决定的关系量。"], y=5.75, h=.7, font=18)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "9. C0/C1：候选关系提升了可预测性")
    add_table(slide, ["实验", "输入", "Spearman", "Pair Acc.", "Per-index", "判断"], [
        ["C0", "O2+Top20六统计", "0.0791", "0.5406", "0.1963", "实例差异显著改善"],
        ["C1", "static prior+C0 residual", "0.1495", "0.5775", "0.2186", "略超static 0.1433/0.5744"],
    ], font=14)
    add_bullets(slide, ["关键转折：candidate-relative utility信号存在；但预测相关性仍需通过检索因果闭环验证。"], y=5.65, h=.8, font=17)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "10. C2-C6：手工门控、删除与局部匹配")
    add_table(slide, ["实验", "MSLS-val结果(hits@1/5/10/20)", "选择", "结论"], [
        ["E0", "685/712/717/719", "-", "基线"],
        ["C2 soft gate", "685/712/717/719", "β=0", "关闭gate最佳"],
        ["C3 hard removal", "684/714/715/720", "remove2", "R@1下降"],
        ["C4 cross-slot", "685/713/715/719", "α=.2", "R@1持平，R@10下降"],
        ["C5 same-slot", "685/712/717/719", "α=0", "局部分数关闭"],
        ["C6 Chamfer/MNN", "685/712/717/719", "α=0", "复杂匹配无增益"],
    ], font=12)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "11. C7-C10：学习式候选重排与路由")
    add_table(slide, ["实验", "做法", "MSLS-val", "结论"], [
        ["C7", "GSV 7特征RankNet", "683/713/715/719", "GSV Top20覆盖99.88%，负迁移"],
        ["C8", "CPH↔SF, 1 epoch", "685/712/717/719", "近恒等，无明确增益"],
        ["C9", "CPH↔SF, 20 epochs", "683/713/715/719", "充分训练后更差"],
        ["C10", "按Top1-Top2 gap路由", "coverage=0", "不启用RankNet最佳"],
    ], font=13)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "12. 六测试集冻结配置：E0基线")
    baseline = {}
    for r in tests:
        if r["method"] == "E0":
            baseline[r["dataset"]] = r
    add_table(slide, ["数据集", "Q", "R", "R@1", "R@5", "R@10", "R@20"], [
        [d, r["queries"], r["references"], f'{r["r@1_pct"]:.4f}', f'{r["r@5_pct"]:.4f}',
         f'{r["r@10_pct"]:.4f}', f'{r["r@20_pct"]:.4f}']
        for d, r in baseline.items()
    ], font=11)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "13. 六测试集ΔR@1：没有稳定赢家")
    methods = ["C3_remove2", "C4_crossmax_a0.2", "C7_gsv_ranknet", "C8_crosscity_e1", "C9_crosscity_e20"]
    datasets = list(baseline)
    lookup = {(r["dataset"], r["method"]): r for r in tests}
    table_rows = []
    for d in datasets:
        table_rows.append([d] + [f'{lookup[(d,m)]["delta_r@1_pp"]:+.4f}' for m in methods])
    table_rows.append(["Macro", "-0.1721", "-0.0269", "-0.0631", "-0.0024", "-0.0511"])
    add_table(slide, ["数据集", "C3", "C4", "C7", "C8", "C9"], table_rows, font=12)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "14. 跨域结果如何解释？")
    add_bullets(slide, [
        "C4：SPED +0.6590 pp，但Nordland −0.3805、AmsterTime −0.4062、SVOX −0.0483 pp。",
        "C7：SPED +0.4942 pp，但Nordland −0.3624、AmsterTime −0.4874 pp。",
        "C3：0胜/1平/5负，直接否定“低分query可安全删除”。",
        "C8最接近E0（macro −0.0024 pp），但1 epoch residual很小，更多是近恒等而非有效提升。",
        "结论：问题不是只发生在MSLS-val；当前utility→descriptor/ranking映射具有域依赖。",
    ], font=19)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "15. 已排除的解释与当前核心瓶颈")
    add_table(slide, ["已排除/证实", "证据", "判断"], [
        ["QRL是否塌缩", "std 0.125，Spearman 0.4415", "未塌缩"],
        ["gate是否真的参与", "gate 0.816~1.199；E2同权重gateoff", "参与但无稳定收益"],
        ["贡献差异是否QRL制造", "E0也有强LOO oracle", "原生BoQ本来就存在"],
        ["是否固定slot重要性即可", "sign-switch 89~100%", "必须instance-specific"],
        ["是否简单冗余导致", "图内Spearman接近0", "否"],
        ["核心瓶颈", "困难负样本主导87.7~96.0%", "target/训练目标未直接优化判别性"],
    ], font=12)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "16. 下一步：从冻结后处理转向联合训练")
    add_bullets(slide, [
        "主方向：L = L_global + λ·L_query/candidate，让O2和FC本身适应hard-negative discrimination。",
        "target：leave-one-query-out descriptor margin、正负样本排名变化，或repeatability−confusability联合目标。",
        "数据：GSV/MSLS-train训练；MSLS-val选择λ/epoch；不得再用本轮六测试集调主配置。",
        "对照：E0/global-only、query-loss、同权重gate-off；至少3 seeds。",
        "门槛：MSLS-val R@1不降、R@5不降、≥2/3 seeds非负，然后冻结一次新held-out评估。",
    ], font=19)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "17. 如果必须使用六测试集选参数")
    add_bullets(slide, [
        "必须改名为跨域开发集；结果标记 TEST-INFORMED optimistic upper bound。",
        "优先保守路由：C4 alpha∈{.01,.02,.05}，coverage∈{5%,10%,20%}。",
        "选择目标：宏平均ΔR@1≥0；最差域≥−0.10pp；负向域≤2；MSLS-val不降。",
        "采用leave-one-dataset-out检查参数稳定性；最终另找未参与选择的数据或预留split。",
    ], font=19)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "18. 汇报结论")
    add_bullets(slide, [
        "1）64个O2 query存在强烈、随图像变化的检索贡献差异。",
        "2）跨视角repeatability可学习，但不是retrieval utility的充分条件。",
        "3）加入候选关系后，C1预测略超static，证明instance-specific信号存在。",
        "4）C2-C10没有把该信号稳定转化为MSLS或六测试集Recall提升。",
        "5）下一步应联合训练hard-negative判别目标，而不是继续堆叠冻结后处理规则。",
    ], y=1.5, font=22)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "附录：证据与产物路径")
    add_bullets(slide, [
        "完整MD：reports/2026-07-31_QRL-BoQ组会汇报/QRL-BoQ_全部实验组会汇报.md",
        "实验总表：QRL-BoQ_实验总表.csv",
        "六测试集：QRL-BoQ_六测试集详细结果.csv",
        "工作簿：QRL-BoQ_组会数据汇总.xlsx",
        "原始测试：/root/qrl_c0_c10_test_diagnostic/all_results.csv",
        "核心代码：src/query_reliability.py、src/boq.py、src/model.py、scripts/train_u0_query_utility.py、scripts/train_candidate_ranknet.py",
    ], font=17)

    prs.save(OUT / "QRL-BoQ_全部实验组会汇报.pptx")


def build_xlsx(rows, tests):
    path = OUT / "QRL-BoQ_组会数据汇总.xlsx"
    wb = xlsxwriter.Workbook(path)
    header = wb.add_format({"bold": True, "bg_color": "#184E77", "font_color": "white", "border": 1})
    cell = wb.add_format({"text_wrap": True, "valign": "top", "border": 1})
    number = wb.add_format({"num_format": "0.0000", "border": 1})

    ws = wb.add_worksheet("实验总表")
    headers = list(rows[0])
    for c, h in enumerate(headers):
        ws.write(0, c, h, header)
    for r, row in enumerate(rows, 1):
        for c, h in enumerate(headers):
            ws.write(r, c, row[h], cell)
    ws.freeze_panes(1, 0)
    ws.set_column(0, 2, 18)
    ws.set_column(3, 7, 42)

    ws = wb.add_worksheet("六测试集详细结果")
    headers = list(tests[0])
    for c, h in enumerate(headers):
        ws.write(0, c, h, header)
    for r, row in enumerate(tests, 1):
        for c, h in enumerate(headers):
            value = row[h]
            if h.startswith("r@") or h.startswith("delta_"):
                ws.write_number(r, c, float(value), number)
            else:
                ws.write(r, c, value, cell)
    ws.freeze_panes(1, 2)
    ws.set_column(0, 1, 22)
    ws.set_column(2, len(headers) - 1, 13)

    ws = wb.add_worksheet("E矩阵_R1")
    headers = ["Dataset", "E0", "E1", "E2", "E2-gateoff", "E3"]
    for c, h in enumerate(headers):
        ws.write(0, c, h, header)
    for r, row in enumerate(E_MATRIX, 1):
        ws.write(r, 0, row[0], cell)
        for c, value in enumerate(row[1:], 1):
            ws.write_number(r, c, value, number)
    ws.set_column(0, 0, 20)
    ws.set_column(1, 5, 15)
    wb.close()


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    rows = experiment_rows()
    tests = test_rows()
    write_csv(OUT / "QRL-BoQ_实验总表.csv", rows, list(rows[0]))
    write_csv(OUT / "QRL-BoQ_六测试集详细结果.csv", tests, list(tests[0]))
    matrix_rows = [
        dict(zip(["Dataset", "E0", "E1", "E2", "E2-gateoff", "E3"], row))
        for row in E_MATRIX
    ]
    write_csv(OUT / "QRL-BoQ_E矩阵_R1.csv", matrix_rows, list(matrix_rows[0]))
    build_xlsx(rows, tests)
    build_ppt(rows, tests)
    manifest = {
        "experiment_rows": len(rows),
        "test_rows": len(tests),
        "e_matrix_rows": len(E_MATRIX),
        "sources": {
            "six_test": "/root/qrl_c0_c10_test_diagnostic/all_results.csv",
            "e_matrix": "/root/qrl_e_matrix/analysis/qrl_e_matrix_results.csv",
            "checkpoint_sha256": "ec3431ae1728ae708b449c06dc016fcf6ce55ae852083911e88ac3a4dfb37d2e",
        },
        "files": sorted(p.name for p in OUT.iterdir()),
    }
    (OUT / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
